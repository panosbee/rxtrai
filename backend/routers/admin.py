"""
RxTrail — Admin router: διαχείριση χρηστών ΜΟΝΟ από τον admin.

Ο admin (settings.admin_email ή user.is_admin):
  - βλέπει όλους τους χρήστες (GET /admin/users)
  - δημιουργεί χρήστες με email + password (POST /admin/users)
  - αλλάζει/επαναφέρει κωδικό (POST /admin/users/{id}/reset-password)
  - ενεργοποιεί/απενεργοποιεί (POST /admin/users/{id}/toggle-active)
  - διαγράφει χρήστη + τα δεδομένα του (DELETE /admin/users/{id})

Οι απλοί χρήστες δεν βλέπουν τίποτα από αυτά — μόνο αλλαγή κωδικού (auth).
"""
import io
import logging
from datetime import datetime, timezone

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field

from backend.auth.deps import get_current_user
from backend.auth.security import hash_password
from backend.config import settings
from backend.db import get_db

logger = logging.getLogger("rxtrail.admin")

router = APIRouter(prefix="/admin", tags=["admin"])

# Οργανισμοί-scoped collections που διαγράφονται μαζί με τον χρήστη
ORG_DATA_COLLECTIONS = [
    "cases", "doctors", "structures", "skus", "sources", "contacts",
    "patients", "prescriptions", "appointments", "fulfilments", "visits",
    "documents", "notifications", "agent_memory", "knowledge_chunks",
    "knowledge_files", "knowledge_lookup", "import_jobs", "agent_procedures",
    "share_files", "share_folders",
]


def _is_admin(user: dict) -> bool:
    return bool(user.get("is_admin")) or (user.get("email") or "").lower() == settings.admin_email.lower()


async def require_admin(current_user: dict = Depends(get_current_user)) -> dict:
    if not _is_admin(current_user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin only")
    return current_user


# ── Schemas ─────────────────────────────────────────────────────────────────

class CreateUserRequest(BaseModel):
    email: EmailStr
    password: str = Field(..., min_length=8, max_length=128)
    name: str = Field(..., min_length=1, max_length=120)


class ResetPasswordRequest(BaseModel):
    password: str = Field(..., min_length=8, max_length=128)


def _user_out(u: dict, org: dict | None) -> dict:
    return {
        "id": str(u["_id"]),
        "email": u.get("email", ""),
        "name": u.get("name", ""),
        "role": u.get("role", "owner"),
        "is_admin": bool(u.get("is_admin")) or (u.get("email") or "").lower() == settings.admin_email.lower(),
        "is_active": bool(u.get("is_active", False)),
        "must_change_password": bool(u.get("must_change_password", False)),
        "org_id": u.get("org_id", ""),
        "org_name": (org or {}).get("name", ""),
        "created_at": u.get("created_at"),
    }


# ── Routes ──────────────────────────────────────────────────────────────────


@router.get("/users")
async def list_users(admin: dict = Depends(require_admin)):
    """Λίστα όλων των χρηστών + το org τους + τα ΔΙΚΑ ΤΟΥΣ δεδομένα (counts).

    Τα counts είναι org-scoped: δείχνουν πόσα δεδομένα ανήκουν στον κάθε
    χρήστη — ο κάθε πωλητής βλέπει/έχει μόνο τα δικά του.
    """
    db = await get_db()
    users = await db.users.find({}).sort("created_at", -1).to_list(length=500)
    out = []
    for u in users:
        org = None
        if u.get("org_id"):
            try:
                org = await db.orgs.find_one({"_id": ObjectId(u["org_id"])})
            except Exception:
                org = None
        counts = {}
        if u.get("org_id"):
            for coll_name, key in [
                ("cases", "cases"), ("doctors", "doctors"),
                ("structures", "structures"), ("skus", "skus"),
                ("sources", "sources"), ("prescriptions", "prescriptions"),
            ]:
                coll = getattr(db, coll_name, None)
                if coll is not None:
                    try:
                        counts[key] = await coll.count_documents({"org_id": u["org_id"]})
                    except Exception:
                        counts[key] = 0
        item = _user_out(u, org)
        item["counts"] = counts
        out.append(item)
    return {"items": out, "total": len(out)}


@router.get("/stats")
async def admin_stats(admin: dict = Depends(require_admin)):
    """Σύνοψη για το admin panel."""
    db = await get_db()
    return {
        "users": await db.users.count_documents({}),
        "orgs": await db.orgs.count_documents({}),
        "cases": await db.cases.count_documents({}),
        "doctors": await db.doctors.count_documents({}),
        "structures": await db.structures.count_documents({}),
        "skus": await db.skus.count_documents({}),
        "sources": await db.sources.count_documents({}),
        "prescriptions": await db.prescriptions.count_documents({}),
        "appointments": await db.appointments.count_documents({}),
        "knowledge_chunks": await db.knowledge_chunks.count_documents({}),
        "knowledge_files": await db.knowledge_files.count_documents({}),
    }


@router.get("/users/{user_id}/stats")
async def user_performance_stats(user_id: str, admin: dict = Depends(require_admin)):
    """Analytics ΑΠΟΔΟΣΗΣ ενός χρήστη — για τα charts του admin.

    Επιστρέφει:
      - counts: σύνολα ανά collection
      - cases_by_status: κατανομή καταστάσεων
      - months + series: 6μηνη χρονοσειρά (cases, γνωματεύσεις, εκτελέσεις, επισκέψεις, τζίρος)
      - totals: αξία ενεργών, συνολικός τζίρος
    """
    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")
    org_id = user.get("org_id", "")

    # ── Σύνολα ──
    counts = {}
    for coll_name, key in [
        ("cases", "cases"), ("doctors", "doctors"), ("structures", "structures"),
        ("skus", "skus"), ("sources", "sources"), ("prescriptions", "prescriptions"),
        ("appointments", "appointments"), ("visits", "visits"),
        ("fulfilments", "fulfilments"), ("documents", "documents"),
    ]:
        coll = getattr(db, coll_name, None)
        counts[key] = await coll.count_documents({"org_id": org_id}) if coll is not None else 0

    # ── Κατανομή καταστάσεων ──
    cases_by_status: dict[str, int] = {}
    async for c in db.cases.aggregate([
        {"$match": {"org_id": org_id}},
        {"$group": {"_id": "$status", "n": {"$sum": 1}}},
    ]):
        cases_by_status[c["_id"]] = c["n"]

    # ── 6μηνη χρονοσειρά ──
    now = datetime.now(timezone.utc)
    months: list[str] = []
    y, m = now.year, now.month
    for _ in range(6):
        months.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    months.reverse()
    start = datetime.strptime(months[0] + "-01", "%Y-%m-%d").replace(tzinfo=timezone.utc)

    async def monthly(collection, date_field: str) -> list[int]:
        buckets = {k: 0 for k in months}
        coll = getattr(db, collection, None)
        if coll is None:
            return [0] * len(months)
        async for row in coll.aggregate([
            {"$match": {"org_id": org_id, date_field: {"$gte": start}}},
            {"$group": {"_id": {"$dateToString": {"format": "%Y-%m", "date": f"${date_field}"}},
                        "n": {"$sum": 1}}},
        ]):
            if row["_id"] in buckets:
                buckets[row["_id"]] = row["n"]
        return [buckets[k] for k in months]

    cases_series = await monthly("cases", "created_at")
    rx_series = await monthly("prescriptions", "created_at")
    fulf_series = await monthly("fulfilments", "fulfilled_at")
    visits_series = await monthly("visits", "at")
    docs_series = await monthly("documents", "created_at")

    # ── Τζίρος ανά μήνα (από εκτελέσεις) + αξία ενεργών ──
    revenue_buckets = {k: 0.0 for k in months}
    async for row in db.fulfilments.aggregate([
        {"$match": {"org_id": org_id, "fulfilled_at": {"$gte": start}}},
        {"$group": {"_id": {"$dateToString": {"format": "%Y-%m", "date": "$fulfilled_at"}},
                    "total": {"$sum": "$total_revenue"}}},
    ]):
        if row["_id"] in revenue_buckets:
            revenue_buckets[row["_id"]] = float(row.get("total") or 0)
    revenue_series = [round(revenue_buckets[k], 2) for k in months]

    active_value = 0.0
    async for c in db.cases.find({"org_id": org_id, "status": {"$in": ["ACTIVE", "AT_RISK"]}},
                                 {"est_monthly_value": 1}):
        try:
            active_value += float(c.get("est_monthly_value") or 0)
        except (TypeError, ValueError):
            pass

    total_revenue = 0.0
    async for row in db.fulfilments.aggregate([
        {"$match": {"org_id": org_id}},
        {"$group": {"_id": None, "total": {"$sum": "$total_revenue"}}},
    ]):
        total_revenue = float(row.get("total") or 0)

    org = None
    if org_id:
        try:
            org = await db.orgs.find_one({"_id": ObjectId(org_id)}, {"name": 1})
        except Exception:
            org = None

    return {
        "user": {
            "id": str(user["_id"]),
            "email": user.get("email", ""),
            "name": user.get("name", ""),
            "created_at": user.get("created_at"),
            "last_login_at": user.get("last_login_at"),
            "is_active": bool(user.get("is_active", False)),
            "org_name": (org or {}).get("name", ""),
        },
        "counts": counts,
        "cases_by_status": cases_by_status,
        "months": months,
        "series": {
            "cases": cases_series,
            "prescriptions": rx_series,
            "fulfilments": fulf_series,
            "visits": visits_series,
            "documents": docs_series,
            "revenue": revenue_series,
        },
        "totals": {
            "active_value": round(active_value, 2),
            "total_revenue": round(total_revenue, 2),
        },
    }


@router.get("/users/{user_id}/export.xlsx")
async def export_user_data_xlsx(user_id: str, admin: dict = Depends(require_admin)):
    """Εξαγωγή δεδομένων ΕΝΟΣ χρήστη σε Excel: Πηγές, Δομές, Περιστατικά, Υγειονομικό.

    Ο admin κατεβάζει ξεχωριστό αρχείο ανά χρήστη, με καθαρά δομημένα φύλλα.
    """
    import io
    from datetime import datetime as _dt

    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")
    org_id = user.get("org_id", "")

    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    def fmt(dt):
        if hasattr(dt, "strftime"):
            return dt.strftime("%Y-%m-%d")
        return str(dt or "")

    wb = Workbook()
    wb.remove(wb.active)

    # ── Σύνοψη ──
    ws = wb.create_sheet(title="Σύνοψη")
    ws.append(["Χρήστης", user.get("name", "")])
    ws.append(["Email", user.get("email", "")])
    ws.append(["Οργανισμός", (await db.orgs.find_one({"_id": ObjectId(org_id)}) or {}).get("name", "") if org_id else ""])
    ws.append(["Εγγραφή", fmt(user.get("created_at"))])
    ws.append(["Ενεργός", "Ναι" if user.get("is_active") else "Όχι"])
    ws.append([])
    ws.append(["Πηγές", await db.sources.count_documents({"org_id": org_id})])
    ws.append(["Δομές", await db.structures.count_documents({"org_id": org_id})])
    ws.append(["Περιστατικά", await db.cases.count_documents({"org_id": org_id})])
    ws.append(["Υγειονομικό προσωπικό", await db.doctors.count_documents({"org_id": org_id})])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 40

    # ── Πηγές ──
    ws = wb.create_sheet(title="Πηγές")
    ws.append(["Όνομα", "Τύπος", "Τηλέφωνο", "Σημειώσεις", " Περιστατικά που έφερε", "Δημιουργήθηκε"])
    async for s in db.sources.find({"org_id": org_id}).sort("name", 1):
        ws.append([
            s.get("name", ""), s.get("type", ""), s.get("phone", ""),
            s.get("notes", ""), s.get("cases_brought", 0), fmt(s.get("created_at")),
        ])

    # ── Δομές ──
    ws = wb.create_sheet(title="Δομές")
    ws.append(["Όνομα", "Τύπος", "Πόλη", "Νομός", "Διεύθυνση", "Τηλέφωνο", "Σημειώσεις"])
    async for st in db.structures.find({"org_id": org_id}).sort("name", 1):
        ws.append([
            st.get("name", ""), st.get("type", ""), st.get("city", ""),
            st.get("prefecture", ""), st.get("address", ""), st.get("phone", ""),
            st.get("notes", ""),
        ])

    # ── Περιστατικά ──
    ws = wb.create_sheet(title="Περιστατικά")
    ws.append(["Ψευδώνυμο", "ΑΜΚΑ", "Κατάσταση", "Περιοχή", "Κατηγορίες", "Αξία €",
               "Γιατρός", "Πηγή", "Κύκλος", "Επόμενη επανάληψη", "Δημιουργήθηκε"])
    for c in await db.cases.find({"org_id": org_id}).sort("created_at", -1).to_list(length=5000):
        doctor_name = ""
        if c.get("doctor_id"):
            try:
                d = await db.doctors.find_one({"_id": ObjectId(c["doctor_id"]), "org_id": org_id}, {"name": 1})
                doctor_name = (d or {}).get("name", "")
            except Exception:
                pass
        source_name = ""
        if c.get("source_id"):
            try:
                sr = await db.sources.find_one({"_id": ObjectId(c["source_id"]), "org_id": org_id}, {"name": 1})
                source_name = (sr or {}).get("name", "")
            except Exception:
                pass
        ws.append([
            c.get("alias", ""), c.get("amka", ""), c.get("status", ""), c.get("area", ""),
            ", ".join(c.get("need_categories") or []), c.get("est_monthly_value", ""),
            doctor_name, source_name, c.get("cycle_number", 0),
            fmt(c.get("next_repeat_due")), fmt(c.get("created_at")),
        ])

    # ── Υγειονομικό Προσωπικό ──
    struct_names: dict[str, str] = {}
    async for st in db.structures.find({"org_id": org_id}, {"name": 1}):
        struct_names[str(st["_id"])] = st.get("name", "")
    ws = wb.create_sheet(title="Υγειονομικό Προσωπικό")
    ws.append(["Όνομα", "Ειδικότητα", "Τύπος προσωπικού", "Χώρος εργασίας", "Ρόλος",
               "Περιοχή", "Διεύθυνση", "Τηλέφωνο", "Email", "Δομές", "Σημειώσεις",
               "Rx 90 ημ.", "Ανοιχτά περιστατικά", "Τελευταία επίσκεψη", "Δημιουργήθηκε"])
    async for d in db.doctors.find({"org_id": org_id}).sort("name", 1):
        structs = ", ".join(struct_names.get(sid, "") for sid in (d.get("structure_ids") or []))
        ws.append([
            d.get("name", ""), d.get("specialty", ""), d.get("personnel_type", ""),
            d.get("workplace_type", ""), d.get("role_title", ""), d.get("area", ""),
            d.get("address", ""), d.get("phone", ""), d.get("email", ""), structs,
            d.get("notes", ""), d.get("rx_count_90d", 0), d.get("open_cases_count", 0),
            fmt(d.get("last_visit_at")), fmt(d.get("created_at")),
        ])

    # Πλάτη στηλών
    for sheet in wb.worksheets:
        widths = {"A": 28, "B": 18, "C": 16, "D": 18, "E": 28, "F": 14, "G": 22}
        for col, w in widths.items():
            sheet.column_dimensions[col].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    stamp = _dt.now().strftime("%Y%m%d_%H%M")
    safe_email = (user.get("email") or "user").replace("@", "_").replace(".", "_")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=rxtrail_{safe_email}_{stamp}.xlsx"},
    )


@router.get("/users/{user_id}/week.xlsx")
async def export_user_week_xlsx(
    user_id: str,
    admin: dict = Depends(require_admin),
    week: str = Query("", description="Οποιαδήποτε ημερομηνία της εβδομάδας (YYYY-MM-DD). Κενό = τρέχουσα."),
):
    """Εβδομαδιαίο πρόγραμμα πωλητή σε Excel — όλα τα ραντεβού της εβδομάδας.

    Ο admin διαλέγει εβδομάδα (?week=YYYY-MM-DD) και κατεβάζει το πρόγραμμα
    του συγκεκριμένου χρήστη, με μέρες, ώρες, γιατρούς, τοποθεσίες και σημειώσεις.
    """
    import io
    from datetime import datetime as _dt, timedelta as _td

    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")
    org_id = user.get("org_id", "")

    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    # Εβδομάδα: Δευτέρα–Κυριακή
    base = datetime.now(timezone.utc)
    if week:
        try:
            base = _dt.fromisoformat(week)
            if base.tzinfo is None:
                base = base.replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    monday = (base - _td(days=base.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    sunday = monday + _td(days=7)

    DAY_NAMES = ["Δευτέρα", "Τρίτη", "Τετάρτη", "Πέμπτη", "Παρασκευή", "Σάββατο", "Κυριακή"]
    KIND_LABELS = {"DOCTOR": "Γιατρός", "CASE": "Περιστατικό", "PERSONAL": "Προσωπικό", "OTHER": "Άλλο"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Εβδομαδιαίο Πρόγραμμα"

    # Κεφαλίδα με στοιχεία
    ws.append([f"Πρόγραμμα: {user.get('name') or user.get('email')}"])
    ws.append([f"Εβδομάδα: {monday.strftime('%d/%m/%Y')} – {(sunday - _td(days=1)).strftime('%d/%m/%Y')}"])
    ws.append([])
    ws.append(["Ημέρα", "Ημερομηνία", "Ώρα", "Διάρκεια (λεπτά)", "Τίτλος", "Τύπος", "Γιατρός", "Τοποθεσία", "Σημειώσεις"])
    header_row = ws.max_row

    # Γιατροί map
    doctor_names: dict[str, str] = {}
    async for d in db.doctors.find({"org_id": org_id}, {"name": 1}):
        doctor_names[str(d["_id"])] = d.get("name", "")

    n = 0
    async for a in db.appointments.find({
        "org_id": org_id,
        "at": {"$gte": monday, "$lt": sunday},
    }).sort("at", 1):
        at = a.get("at")
        if isinstance(at, str):
            try:
                at = _dt.fromisoformat(at)
            except ValueError:
                continue
        if not at:
            continue
        end_at = a.get("end_at")
        duration = ""
        try:
            if end_at:
                if isinstance(end_at, str):
                    end_at = _dt.fromisoformat(end_at)
                duration = int((end_at - at).total_seconds() // 60)
        except Exception:
            duration = ""

        doctor = doctor_names.get(str(a.get("doctor_id") or ""), "")
        ws.append([
            DAY_NAMES[at.weekday()],
            at.strftime("%d/%m/%Y"),
            at.strftime("%H:%M"),
            duration,
            a.get("title", ""),
            KIND_LABELS.get(a.get("kind", ""), a.get("kind", "")),
            doctor,
            a.get("location", ""),
            (a.get("notes") or "").replace("\n", " "),
        ])
        n += 1

    if n == 0:
        ws.append(["—", "", "", "", "Δεν υπάρχουν ραντεβού αυτή την εβδομάδα", "", "", "", ""])

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 13
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 24
    ws.column_dimensions["I"].width = 50

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_email = (user.get("email") or "user").replace("@", "_").replace(".", "_")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=programma_{safe_email}_{monday.strftime('%Y%m%d')}.xlsx"},
    )


@router.get("/export.xlsx")
async def export_all_data_xlsx(admin: dict = Depends(require_admin)):
    """Ο admin κατεβάζει ΟΛΑ τα δεδομένα ΟΛΩΝ των χρηστών σε ένα Excel.

    Κάθε sheet = μία συλλογή. Στήλη «Οργανισμός» δείχνει σε ποιον χρήστη
    ανήκει κάθε γραμμή.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    db = await get_db()

    # org id → name
    org_names: dict[str, str] = {}
    async for o in db.orgs.find({}, {"name": 1}):
        org_names[str(o["_id"])] = o.get("name", "") or "?"

    def org_label(org_id) -> str:
        if not org_id:
            return ""
        return org_names.get(str(org_id), str(org_id))

    # Σειρά φύλλων: (όνομα φύλλου, collection, επιλογή πεδίων ή None=αυτόματα)
    sheets_spec = [
        ("Χρήστες", "users", {"email": 1, "name": 1, "role": 1, "is_admin": 1,
                               "is_active": 1, "must_change_password": 1,
                               "org_id": 1, "created_at": 1, "last_login_at": 1}),
        ("Περιστατικά", "cases", None),
        ("Γιατροί", "doctors", None),
        ("Δομές", "structures", None),
        ("SKU", "skus", None),
        ("Πηγές", "sources", None),
        ("Επαφές", "contacts", None),
        ("Ασθενείς", "patients", None),
        ("Συνταγές", "prescriptions", None),
        ("Ραντεβού", "appointments", None),
        ("Επισκέψεις", "visits", None),
        ("Εκτελέσεις", "fulfilments", None),
        ("Αρχεία Γνώσης", "knowledge_files", None),
    ]

    EXCLUDED_FIELDS = {"_id", "embedding", "name_embedding", "password_hash", "vector"}
    MAX_ROWS_PER_SHEET = 5000

    def _cell(v):
        if v is None:
            return ""
        if isinstance(v, ObjectId):
            return str(v)
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M")
        if isinstance(v, (dict, list)):
            import json
            s = json.dumps(v, ensure_ascii=False, default=str)
            return s[:2000]
        return v

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, coll_name, projection in sheets_spec:
        coll = getattr(db, coll_name, None)
        if coll is None:
            continue
        rows = []
        cursor = coll.find({}, projection)
        if projection is None:
            cursor = cursor.limit(MAX_ROWS_PER_SHEET + 1)
        async for doc in cursor:
            rows.append(doc)
            if len(rows) > MAX_ROWS_PER_SHEET:
                break

        if not rows:
            continue

        # Δυναμικά headers από τα πρώτα docs
        headers: list[str] = []
        for doc in rows[:50]:
            for k in doc.keys():
                if k in EXCLUDED_FIELDS or k in headers:
                    continue
                headers.append(k)

        ws = wb.create_sheet(title=sheet_name[:31])
        ws.append(["Οργανισμός"] + headers)

        for doc in rows:
            ws.append([org_label(doc.get("org_id"))] + [_cell(doc.get(h)) for h in headers])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=rxtrail_all_data_{stamp}.xlsx"},
    )


@router.post("/users", status_code=status.HTTP_201_CREATED)
async def create_user(req: CreateUserRequest, admin: dict = Depends(require_admin)):
    """Ο admin δημιουργεί χρήστη: email + αρχικός κωδικός + δικό του org.

    Ο χρήστης μπαίνει με must_change_password=True — αναγκάζεται να αλλάξει
    κωδικό στην πρώτη είσοδο.
    """
    db = await get_db()
    email = req.email.lower()

    existing = await db.users.find_one({"email": email})
    if existing:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email already registered")

    now = datetime.now(timezone.utc)
    org_result = await db.orgs.insert_one({
        "name": req.name,
        "settings": {
            "repeat_notify_day": 25,
            "at_risk_day": 35,
            "new_sla_days": 5,
            "assessed_sla_days": 7,
            "prescribed_sla_days": 7,
            "photo_retention_days": 30,
        },
        "created_at": now,
    })
    org_id = str(org_result.inserted_id)

    user_result = await db.users.insert_one({
        "email": email,
        "password_hash": hash_password(req.password),
        "name": req.name,
        "org_id": org_id,
        "role": "owner",
        "is_admin": False,
        "is_active": True,
        "must_change_password": True,
        "created_at": now,
        "updated_at": now,
    })

    logger.info(f"Admin {admin.get('email')} created user {email} (org={org_id})")
    return {
        "id": str(user_result.inserted_id),
        "email": email,
        "name": req.name,
        "org_id": org_id,
        "must_change_password": True,
        "message": f"Ο χρήστης {email} δημιουργήθηκε. Θα αλλάξει κωδικό στην πρώτη είσοδο.",
    }


@router.post("/users/{user_id}/reset-password")
async def reset_password(user_id: str, req: ResetPasswordRequest, admin: dict = Depends(require_admin)):
    """Ο admin ορίζει νέο κωδικό για τον χρήστη."""
    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")

    await db.users.update_one(
        {"_id": ObjectId(user_id)},
        {"$set": {
            "password_hash": hash_password(req.password),
            "must_change_password": True,
            "updated_at": datetime.now(timezone.utc),
        }},
    )
    logger.info(f"Admin {admin.get('email')} reset password for {user.get('email')}")
    return {"ok": True, "message": f"Νέος κωδικός για {user.get('email')}. Θα αλλάξει στην επόμενη είσοδο."}


@router.post("/users/{user_id}/toggle-active")
async def toggle_active(user_id: str, admin: dict = Depends(require_admin)):
    """Ενεργοποίηση / απενεργοποίηση λογαριασμού."""
    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")
    if str(user["_id"]) == str(admin["_id"]):
        raise HTTPException(400, "Δεν μπορείς να απενεργοποιήσεις τον εαυτό σου")

    new_state = not bool(user.get("is_active", False))
    await db.users.update_one(
        {"_id": ObjectId(user_id)},
        {"$set": {"is_active": new_state, "updated_at": datetime.now(timezone.utc)}},
    )
    logger.info(f"Admin {admin.get('email')} set {user.get('email')} active={new_state}")
    return {"ok": True, "is_active": new_state, "message": f"Ο λογαριασμός {'ενεργοποιήθηκε' if new_state else 'απενεργοποιήθηκε'}."}


@router.delete("/users/{user_id}")
async def delete_user(user_id: str, admin: dict = Depends(require_admin)):
    """Διαγραφή χρήστη + του org του + όλων των δεδομένων του org."""
    db = await get_db()
    if not ObjectId.is_valid(user_id):
        raise HTTPException(404, "User not found")
    user = await db.users.find_one({"_id": ObjectId(user_id)})
    if not user:
        raise HTTPException(404, "User not found")
    if str(user["_id"]) == str(admin["_id"]):
        raise HTTPException(400, "Δεν μπορείς να διαγράψεις τον εαυτό σου")

    org_id = user.get("org_id", "")
    deleted_counts = {}
    if org_id:
        for coll_name in ORG_DATA_COLLECTIONS:
            coll = getattr(db, coll_name, None)
            if coll is None:
                continue
            try:
                res = await coll.delete_many({"org_id": org_id})
                deleted_counts[coll_name] = res.deleted_count
            except Exception:
                deleted_counts[coll_name] = -1
        try:
            await db.orgs.delete_one({"_id": ObjectId(org_id)})
            deleted_counts["orgs"] = 1
        except Exception:
            deleted_counts["orgs"] = 0

    await db.users.delete_one({"_id": ObjectId(user_id)})
    logger.warning(f"Admin {admin.get('email')} DELETED user {user.get('email')} (org={org_id})")
    return {
        "ok": True,
        "message": f"Ο χρήστης {user.get('email')} και τα δεδομένα του διαγράφηκαν οριστικά.",
        "deleted": deleted_counts,
    }
